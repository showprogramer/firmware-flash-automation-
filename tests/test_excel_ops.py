import builtins
from pathlib import Path

import openpyxl
import pytest

from handcontrol.core.excel_ops import load_excel_row, load_excel_status, merge_backup_excel, read_all_excel_rows, write_excel_record


def _logs():
    messages = []
    return messages, messages.append


def test_write_excel_record_creates_new_file(tmp_path: Path):
    excel = tmp_path / "records.xlsx"
    logs, log_fn = _logs()

    result = write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.2.3",
        remark="待确认",
        log_fn=log_fn,
    )

    assert result["ok"] is True
    assert result["reason"] == "ok"
    assert result["written_row"]["model"] == "L36"
    assert result["written_row"]["version"] == "V1.2.3"
    assert result["written_row"]["remark"] == "待确认"
    assert result["written_row"]["date"]
    assert excel.exists()
    assert any("Excel 已写入" in msg for msg in logs)

    wb = openpyxl.load_workbook(excel)
    ws = wb["Sheet1"]
    assert ws.cell(row=3, column=2).value == "L36"
    assert ws.cell(row=3, column=6).value == "V1.2.3"
    assert ws.cell(row=3, column=9).value == "待确认"
    wb.close()


def test_write_excel_record_updates_existing_row(tmp_path: Path):
    excel = tmp_path / "records.xlsx"

    first = write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.2.3",
        remark="待确认",
        log_fn=lambda _: None,
    )
    assert first["ok"] is True

    second = write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.2.3",
        remark="测试通过",
        logo="品牌A",
        language="中、英",
        salesman="张三",
        attachment="定制",
        log_fn=lambda _: None,
    )
    assert second["ok"] is True
    assert second["written_row"]["logo"] == "品牌A"
    assert second["written_row"]["salesman"] == "张三"
    assert second["written_row"]["language"] == "中、英"
    assert second["written_row"]["attachment"] == "定制"
    assert second["written_row"]["remark"] == "测试通过"

    wb = openpyxl.load_workbook(excel)
    ws = wb["Sheet1"]
    # still one data row at row 3, should be updated rather than appended
    assert ws.max_row == 3
    assert ws.cell(row=3, column=3).value == "品牌A"
    assert ws.cell(row=3, column=4).value == "张三"
    assert ws.cell(row=3, column=5).value == "中、英"
    assert ws.cell(row=3, column=8).value == "定制"
    assert ws.cell(row=3, column=9).value == "测试通过"
    wb.close()


def test_write_excel_record_returns_locked_and_writes_tmp(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    excel = tmp_path / "records.xlsx"
    excel.write_bytes(b"placeholder")
    logs, log_fn = _logs()

    original_open = builtins.open

    def fake_open(path, mode="r", *args, **kwargs):
        p = Path(path)
        if p == excel and mode == "r+b":
            raise PermissionError("locked")
        return original_open(path, mode, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", fake_open)

    result = write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L50S",
        version="V3.4.5",
        remark="待确认",
        log_fn=log_fn,
    )

    expected_tmp = excel.with_name("records_刷机记录_待导入.xlsx")
    assert result["ok"] is False
    assert result["reason"] == "locked"
    assert result["written_row"] is None
    assert result["tmp_path"] == str(expected_tmp)
    assert expected_tmp.exists()
    assert any("被 WPS/Excel 占用" in msg for msg in logs)


def test_write_excel_record_returns_write_failed_on_exception(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    excel = tmp_path / "records.xlsx"
    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )

    def raise_load(*args, **kwargs):
        raise RuntimeError("mock load failure")

    monkeypatch.setattr("handcontrol.core.excel_ops.openpyxl.load_workbook", raise_load)
    logs, log_fn = _logs()

    result = write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="测试通过",
        log_fn=log_fn,
    )

    assert result["ok"] is False
    assert result["reason"] == "write_failed"
    assert result["written_row"] is None
    assert "mock load failure" in result["error"]
    assert any("Excel 写入失败" in msg for msg in logs)


def test_read_helpers_return_expected_data(tmp_path: Path):
    excel = tmp_path / "records.xlsx"

    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L66",
        version="V9.9.9",
        remark="测试通过",
        logo="品牌B",
        language="英文",
        salesman="李四",
        attachment="可通用",
        log_fn=lambda _: None,
    )

    rows = read_all_excel_rows(str(excel), "Sheet1")
    assert len(rows) == 1
    assert rows[0][1] == "L66"
    assert rows[0][5] == "V9.9.9"
    assert rows[0][7] == "可通用"
    assert rows[0][8] == "测试通过"

    status = load_excel_status(str(excel), "Sheet1")
    assert status[("L66", "V9.9.9")] == "测试通过"

    row = load_excel_row(str(excel), "Sheet1", "L66", "V9.9.9")
    assert row == {
        "logo": "品牌B",
        "salesman": "李四",
        "language": "英文",
        "attachment": "可通用",
        "remark": "测试通过",
    }

def test_serial_numbers_increment_for_new_rows(tmp_path: Path):
    excel = tmp_path / "records.xlsx"

    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )
    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L50S",
        version="V2.0.0",
        remark="测试通过",
        log_fn=lambda _: None,
    )

    wb = openpyxl.load_workbook(excel)
    ws = wb["Sheet1"]
    assert ws.cell(row=3, column=1).value == 1
    assert ws.cell(row=4, column=1).value == 2
    wb.close()


def test_serial_numbers_stay_continuous_after_update(tmp_path: Path):
    excel = tmp_path / "records.xlsx"

    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )
    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L50S",
        version="V2.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )

    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="测试通过",
        logo="品牌A",
        log_fn=lambda _: None,
    )

    wb = openpyxl.load_workbook(excel)
    ws = wb["Sheet1"]
    assert ws.cell(row=3, column=1).value == 1
    assert ws.cell(row=4, column=1).value == 2
    assert ws.cell(row=3, column=9).value == "测试通过"
    wb.close()


def test_serial_numbers_ignore_blank_rows(tmp_path: Path):
    excel = tmp_path / "records.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="现有手控UI明细")
    ws.cell(row=2, column=1, value="序号")
    ws.cell(row=2, column=2, value="型号")
    ws.cell(row=2, column=6, value="版本号")
    ws.cell(row=3, column=2, value="L36")
    ws.cell(row=3, column=6, value="V1.0.0")
    ws.cell(row=4, column=1, value=999)
    ws.cell(row=5, column=2, value="L50S")
    ws.cell(row=5, column=6, value="V2.0.0")
    wb.save(excel)
    wb.close()

    write_excel_record(
        excel_path=str(excel),
        sheet_name="Sheet1",
        model="L50S",
        version="V2.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )

    wb = openpyxl.load_workbook(excel)
    ws = wb["Sheet1"]
    assert ws.cell(row=3, column=1).value == 1
    assert ws.cell(row=4, column=1).value is None
    assert ws.cell(row=5, column=1).value == 2
    wb.close()





def test_merge_backup_excel_appends_only_new_rows_and_deletes_backup(tmp_path: Path):
    main_excel = tmp_path / "records.xlsx"
    backup_excel = tmp_path / "records_刷机记录_待导入.xlsx"

    write_excel_record(
        excel_path=str(main_excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )
    write_excel_record(
        excel_path=str(backup_excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )
    write_excel_record(
        excel_path=str(backup_excel),
        sheet_name="Sheet1",
        model="L50S",
        version="V2.0.0",
        remark="测试通过",
        logo="品牌B",
        attachment="定制",
        log_fn=lambda _: None,
    )

    result = merge_backup_excel(str(main_excel), str(backup_excel), "Sheet1", log_fn=lambda _: None)

    assert result["ok"] is True
    assert result["merged_count"] == 1
    assert result["skipped_count"] == 1
    assert result["source_deleted"] is True
    assert backup_excel.exists() is False

    rows = read_all_excel_rows(str(main_excel), "Sheet1")
    assert len(rows) == 2
    assert rows[1][1] == "L50S"
    assert rows[1][7] == "定制"
    assert rows[1][8] == "测试通过"


def test_merge_backup_excel_returns_no_rows_when_backup_has_no_new_data(tmp_path: Path):
    main_excel = tmp_path / "records.xlsx"
    backup_excel = tmp_path / "records_刷机记录_待导入.xlsx"

    write_excel_record(
        excel_path=str(main_excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="待确认",
        log_fn=lambda _: None,
    )
    write_excel_record(
        excel_path=str(backup_excel),
        sheet_name="Sheet1",
        model="L36",
        version="V1.0.0",
        remark="测试通过",
        log_fn=lambda _: None,
    )

    result = merge_backup_excel(str(main_excel), str(backup_excel), "Sheet1", log_fn=lambda _: None)

    assert result["ok"] is False
    assert result["reason"] == "no_rows"
    assert result["merged_count"] == 0
    assert result["skipped_count"] == 1
    assert backup_excel.exists() is True
