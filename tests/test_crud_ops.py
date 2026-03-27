"""
CRUD 操作测试
=============
阶段1 smoke：搜索过滤逻辑（纯内存，不依赖 UI）
阶段2：双击回填场景 —— 离线历史记录可修改（update_excel_field 全路径）
阶段3：delete_excel_row 全分支（ok / locked / not_found / write_failed）
"""
import builtins
from pathlib import Path

import openpyxl
import pytest

from handcontrol.core.excel_ops import delete_excel_row, update_excel_field, write_excel_record


# ─────────────────────────────────────────────────────────────────────────────
# 共用工具
# ─────────────────────────────────────────────────────────────────────────────

def _logs():
    messages: list[str] = []
    return messages, messages.append


def _make_excel(tmp_path: Path, rows: list[dict]) -> Path:
    """建一个含多行数据的 Excel 文件，返回路径。"""
    excel = tmp_path / "records.xlsx"
    for row in rows:
        write_excel_record(
            excel_path=str(excel),
            sheet_name="Sheet1",
            model=row["model"],
            version=row["version"],
            remark=row.get("remark", "待确认"),
            logo=row.get("logo", ""),
            language=row.get("language", ""),
            salesman=row.get("salesman", ""),
            attachment=row.get("attachment", ""),
            log_fn=lambda _: None,
        )
    return excel


def _preview_row_matches(row: dict, keyword: str) -> bool:
    """
    阶段1：纯内存搜索过滤逻辑（与 app.py _row_matches_keyword 保持一致）。
    单独提取为纯函数，方便单元测试，不依赖 tkinter。
    """
    keyword = keyword.strip().lower()
    if not keyword:
        return True
    fields = ("model", "version", "logo", "salesman", "language", "attachment", "remark", "serial")
    return any(keyword in str(row.get(f, "")).lower() for f in fields)


# ─────────────────────────────────────────────────────────────────────────────
# 阶段1：搜索过滤 smoke 测试（纯内存，不依赖 UI）
# ─────────────────────────────────────────────────────────────────────────────

class TestSearchFilter:
    """验证搜索关键词过滤逻辑覆盖各字段。"""

    ROWS = [
        {"model": "L36",  "version": "V1.0.0", "logo": "中性",  "language": "中、英、越", "attachment": "可通用", "salesman": "张三", "remark": "测试通过",  "serial": "1"},
        {"model": "L50S", "version": "V2.0.0", "logo": "通用",  "language": "英文",      "attachment": "定制",   "salesman": "李四", "remark": "待确认",    "serial": "2"},
        {"model": "L66",  "version": "V3.1.0", "logo": "定制",  "language": "中、英",    "attachment": "可通用", "salesman": "王五", "remark": "测试失败",  "serial": "3"},
    ]

    def test_empty_keyword_matches_all(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "")]
        assert len(result) == 3

    def test_match_by_model(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "L50")]
        assert len(result) == 1
        assert result[0]["model"] == "L50S"

    def test_match_by_version(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "3.1")]
        assert len(result) == 1
        assert result[0]["version"] == "V3.1.0"

    def test_match_by_logo(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "定制")]
        assert any(r["logo"] == "定制" for r in result)

    def test_match_by_salesman(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "李四")]
        assert len(result) == 1
        assert result[0]["salesman"] == "李四"

    def test_match_by_language(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "英文")]
        # "英文" 精确在 L50S 行，但 "中、英、越" 和 "中、英" 也包含"英"字
        # 本测试验证 "英文" 精确命中逻辑
        assert any(r["model"] == "L50S" for r in result)

    def test_match_by_remark(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "测试通过")]
        assert len(result) == 1
        assert result[0]["remark"] == "测试通过"

    def test_match_by_serial(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "2")]
        assert any(r["serial"] == "2" for r in result)

    def test_no_match_returns_empty(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "不存在的关键词xyz")]
        assert result == []

    def test_case_insensitive(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "l36")]
        assert len(result) == 1
        assert result[0]["model"] == "L36"

    def test_whitespace_keyword_matches_all(self):
        result = [r for r in self.ROWS if _preview_row_matches(r, "   ")]
        assert len(result) == 3


# ─────────────────────────────────────────────────────────────────────────────
# 阶段2：update_excel_field（双击回填后离线修改历史记录）
# ─────────────────────────────────────────────────────────────────────────────

class TestUpdateExcelField:
    """验证精确字段更新（logo / salesman / language / remark）的全路径。"""

    def test_update_logo_ok(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0", "logo": "中性"}])
        logs, log_fn = _logs()

        result = update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "logo", "定制", log_fn)

        assert result["ok"] is True
        assert result["reason"] == "ok"
        assert any("已更新" in m for m in logs)
        wb = openpyxl.load_workbook(excel)
        assert wb["Sheet1"].cell(row=3, column=3).value == "定制"
        wb.close()

    def test_update_salesman_ok(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L50S", "version": "V2.0.0", "salesman": "张三"}])
        result = update_excel_field(str(excel), "Sheet1", "L50S", "V2.0.0", "salesman", "李四", lambda _: None)

        assert result["ok"] is True
        wb = openpyxl.load_workbook(excel)
        assert wb["Sheet1"].cell(row=3, column=4).value == "李四"
        wb.close()

    def test_update_language_ok(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L66", "version": "V3.0.0"}])
        result = update_excel_field(str(excel), "Sheet1", "L66", "V3.0.0", "language", "中、英", lambda _: None)

        assert result["ok"] is True
        wb = openpyxl.load_workbook(excel)
        assert wb["Sheet1"].cell(row=3, column=5).value == "中、英"
        wb.close()

    def test_update_remark_ok(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0", "remark": "待确认"}])
        result = update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "remark", "测试通过", lambda _: None)

        assert result["ok"] is True
        wb = openpyxl.load_workbook(excel)
        assert wb["Sheet1"].cell(row=3, column=9).value == "测试通过"
        wb.close()

    def test_update_unknown_field_returns_unknown_field(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        logs, log_fn = _logs()

        result = update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "nonexistent", "x", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "unknown_field"
        assert any("未知字段" in m for m in logs)

    def test_update_not_found_row(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        logs, log_fn = _logs()

        result = update_excel_field(str(excel), "Sheet1", "L99", "V9.9.9", "logo", "定制", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "not_found"
        assert any("未找到" in m for m in logs)

    def test_update_file_not_exist(self, tmp_path: Path):
        excel = tmp_path / "missing.xlsx"
        result = update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "logo", "定制", lambda _: None)

        assert result["ok"] is False
        assert result["reason"] == "not_found"

    def test_update_locked_file(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        original_open = builtins.open

        def fake_open(path, mode="r", *args, **kwargs):
            if Path(path) == excel and mode == "r+b":
                raise PermissionError("locked")
            return original_open(path, mode, *args, **kwargs)

        monkeypatch.setattr(builtins, "open", fake_open)
        logs, log_fn = _logs()

        result = update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "logo", "定制", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "locked"
        assert any("占用" in m for m in logs)

    def test_update_write_failed_on_exception(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])

        def raise_load(*args, **kwargs):
            raise RuntimeError("mock failure")

        monkeypatch.setattr("handcontrol.core.excel_ops.openpyxl.load_workbook", raise_load)
        logs, log_fn = _logs()

        result = update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "logo", "定制", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "write_failed"
        assert "mock failure" in result["error"]

    def test_update_does_not_affect_other_rows(self, tmp_path: Path):
        """修改一行不影响其他行的数据。"""
        excel = _make_excel(tmp_path, [
            {"model": "L36",  "version": "V1.0.0", "logo": "中性"},
            {"model": "L50S", "version": "V2.0.0", "logo": "通用"},
        ])
        update_excel_field(str(excel), "Sheet1", "L36", "V1.0.0", "logo", "定制", lambda _: None)

        wb = openpyxl.load_workbook(excel)
        ws = wb["Sheet1"]
        # L50S 行不受影响
        assert ws.cell(row=4, column=3).value == "通用"
        wb.close()

    def test_update_is_case_insensitive_on_model_version(self, tmp_path: Path):
        """型号/版本号大小写不敏感匹配。"""
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        result = update_excel_field(str(excel), "Sheet1", "l36", "v1.0.0", "logo", "定制", lambda _: None)

        assert result["ok"] is True
        wb = openpyxl.load_workbook(excel)
        assert wb["Sheet1"].cell(row=3, column=3).value == "定制"
        wb.close()


# ─────────────────────────────────────────────────────────────────────────────
# 阶段3：delete_excel_row 全分支测试
# ─────────────────────────────────────────────────────────────────────────────

class TestDeleteExcelRow:
    """delete_excel_row 的 ok / locked / not_found / write_failed 全分支。"""

    def test_delete_ok_single_row(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        logs, log_fn = _logs()

        result = delete_excel_row(str(excel), "Sheet1", "L36", "V1.0.0", log_fn)

        assert result["ok"] is True
        assert result["reason"] == "ok"
        assert any("已删除" in m for m in logs)
        wb = openpyxl.load_workbook(excel)
        ws = wb["Sheet1"]
        # 数据行应该消失，只剩两行标题
        assert ws.max_row == 2
        wb.close()

    def test_delete_ok_serial_renumbered(self, tmp_path: Path):
        """删除中间行后，序号连续。"""
        excel = _make_excel(tmp_path, [
            {"model": "L36",  "version": "V1.0.0"},
            {"model": "L50S", "version": "V2.0.0"},
            {"model": "L66",  "version": "V3.0.0"},
        ])
        delete_excel_row(str(excel), "Sheet1", "L50S", "V2.0.0", lambda _: None)

        wb = openpyxl.load_workbook(excel)
        ws = wb["Sheet1"]
        assert ws.max_row == 4          # 2 标题 + 2 数据
        assert ws.cell(row=3, column=1).value == 1   # L36
        assert ws.cell(row=4, column=1).value == 2   # L66
        assert ws.cell(row=3, column=2).value == "L36"
        assert ws.cell(row=4, column=2).value == "L66"
        wb.close()

    def test_delete_not_found(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        logs, log_fn = _logs()

        result = delete_excel_row(str(excel), "Sheet1", "L99", "V9.9.9", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "not_found"
        assert any("未找到" in m for m in logs)

    def test_delete_file_not_exist(self, tmp_path: Path):
        excel = tmp_path / "missing.xlsx"
        logs, log_fn = _logs()

        result = delete_excel_row(str(excel), "Sheet1", "L36", "V1.0.0", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "not_found"

    def test_delete_locked_file(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        original_open = builtins.open

        def fake_open(path, mode="r", *args, **kwargs):
            if Path(path) == excel and mode == "r+b":
                raise PermissionError("locked")
            return original_open(path, mode, *args, **kwargs)

        monkeypatch.setattr(builtins, "open", fake_open)
        logs, log_fn = _logs()

        result = delete_excel_row(str(excel), "Sheet1", "L36", "V1.0.0", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "locked"
        assert any("占用" in m for m in logs)

    def test_delete_write_failed_on_exception(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])

        def raise_load(*args, **kwargs):
            raise RuntimeError("mock delete failure")

        monkeypatch.setattr("handcontrol.core.excel_ops.openpyxl.load_workbook", raise_load)
        logs, log_fn = _logs()

        result = delete_excel_row(str(excel), "Sheet1", "L36", "V1.0.0", log_fn)

        assert result["ok"] is False
        assert result["reason"] == "write_failed"
        assert "mock delete failure" in result["error"]

    def test_delete_is_case_insensitive(self, tmp_path: Path):
        excel = _make_excel(tmp_path, [{"model": "L36", "version": "V1.0.0"}])
        result = delete_excel_row(str(excel), "Sheet1", "l36", "v1.0.0", lambda _: None)

        assert result["ok"] is True
        wb = openpyxl.load_workbook(excel)
        assert wb["Sheet1"].max_row == 2
        wb.close()

    def test_delete_does_not_affect_other_rows_data(self, tmp_path: Path):
        """删除一行后，其他行数据完整保留。"""
        excel = _make_excel(tmp_path, [
            {"model": "L36",  "version": "V1.0.0", "logo": "中性", "language": "中、英", "salesman": "张三", "remark": "测试通过"},
            {"model": "L50S", "version": "V2.0.0", "logo": "通用", "language": "英文",   "salesman": "李四", "remark": "待确认"},
        ])
        delete_excel_row(str(excel), "Sheet1", "L36", "V1.0.0", lambda _: None)

        wb = openpyxl.load_workbook(excel)
        ws = wb["Sheet1"]
        assert ws.max_row == 3      # 2 标题 + 1 数据
        assert ws.cell(row=3, column=2).value == "L50S"
        assert ws.cell(row=3, column=3).value == "通用"
        assert ws.cell(row=3, column=5).value == "英文"
        assert ws.cell(row=3, column=9).value == "待确认"
        wb.close()






