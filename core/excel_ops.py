from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from core.settings import EXCEL_HEADER_ROW


def load_excel(path: str):
    if not Path(path).exists():
        raise FileNotFoundError(f"找不到 Excel 文件: {path}")
    return openpyxl.load_workbook(path)


def find_or_create_row(ws, model: str, version: str) -> int:
    """
    Find matching model+version row or return next append row.
    Model in col B, version in col F, data starts at row 3.
    """
    for row in ws.iter_rows(min_row=EXCEL_HEADER_ROW + 1, max_row=ws.max_row):
        b_val = str(row[1].value or "").strip()
        f_val = str(row[5].value or "").strip()
        if b_val.upper() == model.upper() and f_val.upper() == version.upper():
            return row[0].row
    return ws.max_row + 1


def _do_write(
    excel_path: str,
    sheet_name: str,
    model: str,
    version: str,
    remark: str,
    rom_path: str,
    log_fn,
    logo: str = "",
    language: str = "",
    salesman: str = "",
    is_tmp: bool = False,
):
    try:
        p = Path(excel_path)
        if p.exists():
            wb = openpyxl.load_workbook(excel_path)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(["现有摩众手控UI明细"])
            ws.append(["序号", "型号", "logo", "业务员", "语言", "版本号", "完成日期", "附图", "备注"])

        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        row_num = find_or_create_row(ws, model, version)

        if rom_path and Path(rom_path).exists():
            mtime = Path(rom_path).stat().st_mtime
            date_str = datetime.fromtimestamp(mtime).strftime("%Y.%m.%d")
        else:
            date_str = datetime.now().strftime("%Y.%m.%d")

        yellow = PatternFill("solid", fgColor="FFFF99")
        green = PatternFill("solid", fgColor="C6EFCE")

        ws.cell(row=row_num, column=2, value=model)
        if logo:
            ws.cell(row=row_num, column=3, value=logo)
        if salesman:
            ws.cell(row=row_num, column=4, value=salesman)
        if language:
            ws.cell(row=row_num, column=5, value=language)
        ws.cell(row=row_num, column=6, value=version)
        ws.cell(row=row_num, column=7, value=date_str)
        ws.cell(row=row_num, column=9, value=remark)

        fill = yellow if remark == "待确认" else (green if remark == "测试通过" else PatternFill())
        for col in range(1, 10):
            ws.cell(row=row_num, column=col).fill = fill

        wb.save(excel_path)
        label = "（备用文件）" if is_tmp else ""
        log_fn(f"  Excel 已写入{label}: 第{row_num}行  {model} {version}  {date_str}  [{remark}]")
        if is_tmp:
            log_fn(f"  备用文件路径: {excel_path}")
        return {
            "ok": True,
            "reason": "ok",
            "tmp_path": "",
            "error": "",
        }
    except Exception as e:
        log_fn(f"  Excel 写入失败: {e}")
        return {
            "ok": False,
            "reason": "write_failed",
            "tmp_path": "",
            "error": str(e),
        }


def write_excel_record(
    excel_path: str,
    sheet_name: str,
    model: str,
    version: str,
    remark: str = "待确认",
    rom_path: str = "",
    logo: str = "",
    language: str = "",
    salesman: str = "",
    log_fn=print,
):
    """
    Write or update a record in Excel.
    Returns dict with keys: ok, reason, tmp_path, error.
    If file is locked by WPS/Excel, write to a temporary backup file.
    """
    excel_path = Path(excel_path)

    def _is_locked(p: Path) -> bool:
        if not p.exists():
            return False
        try:
            with open(p, "r+b"):
                return False
        except FileNotFoundError:
            return False
        except (IOError, PermissionError):
            return True

    if _is_locked(excel_path):
        log_fn("  ⚠ Excel 文件被 WPS/Excel 占用，将保存为备用文件，请手动合并或关闭后重试")
        tmp_path = excel_path.with_name(excel_path.stem + "_刷机记录_待导入.xlsx")
        tmp_result = _do_write(
            str(tmp_path),
            sheet_name,
            model,
            version,
            remark,
            rom_path,
            log_fn,
            logo=logo,
            language=language,
            salesman=salesman,
            is_tmp=True,
        )
        return {
            "ok": False,
            "reason": "locked",
            "tmp_path": str(tmp_path),
            "error": str(tmp_result.get("error", "")),
        }

    return _do_write(
        str(excel_path),
        sheet_name,
        model,
        version,
        remark,
        rom_path,
        log_fn,
        logo=logo,
        language=language,
        salesman=salesman,
    )


def read_all_excel_rows(excel_path: str, sheet_name: str) -> list[list]:
    """Read all data rows (skip first two title/header rows)."""
    rows = []
    p = Path(excel_path)
    if not p.exists():
        return rows
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        for row in ws.iter_rows(min_row=EXCEL_HEADER_ROW + 1, max_row=ws.max_row, values_only=True):
            if any(v is not None for v in row):
                rows.append([str(v or "") for v in row])
        wb.close()
    except Exception:
        pass
    return rows


def load_excel_status(excel_path: str, sheet_name: str) -> dict:
    """Return {(MODEL, VERSION): remark} from Excel."""
    result = {}
    excel = Path(excel_path)
    if not excel.exists():
        return result
    try:
        wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        for row in ws.iter_rows(min_row=EXCEL_HEADER_ROW + 1, max_row=ws.max_row, values_only=True):
            if len(row) >= 9:
                m = str(row[1] or "").strip().upper()
                v = str(row[5] or "").strip().upper()
                r = str(row[8] or "").strip()
                if m and v:
                    result[(m, v)] = r
        wb.close()
    except Exception:
        pass
    return result


def load_excel_row(excel_path: str, sheet_name: str, model: str, version: str) -> dict:
    """Return one row details for model+version, or empty dict."""
    excel = Path(excel_path)
    if not excel.exists():
        return {}
    try:
        wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        for row in ws.iter_rows(min_row=EXCEL_HEADER_ROW + 1, max_row=ws.max_row, values_only=True):
            if str(row[1] or "").strip().upper() == model.upper() and str(row[5] or "").strip().upper() == version.upper():
                wb.close()
                return {
                    "logo": str(row[2] or ""),
                    "salesman": str(row[3] or ""),
                    "language": str(row[4] or ""),
                    "remark": str(row[8] or ""),
                }
        wb.close()
    except Exception:
        pass
    return {}
